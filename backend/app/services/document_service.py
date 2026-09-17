"""
Converts a generated petition draft (plain text) into downloadable
PDF or DOCX files, saved under settings.UPLOAD_DIR.
"""
import os
import uuid

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.core.config import settings


def _ensure_upload_dir() -> str:
    os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
    return settings.UPLOAD_DIR


def petition_to_pdf(title: str, content_text: str) -> str:
    directory = _ensure_upload_dir()
    filename = f"petition_{uuid.uuid4().hex}.pdf"
    path = os.path.join(directory, filename)

    doc = SimpleDocTemplate(
        path, pagesize=A4,
        leftMargin=1 * inch, rightMargin=1 * inch, topMargin=1 * inch, bottomMargin=1 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("PetitionTitle", parent=styles["Heading1"], alignment=1, spaceAfter=18)
    body_style = ParagraphStyle("PetitionBody", parent=styles["BodyText"], fontSize=11, leading=16, spaceAfter=10)

    story = [Paragraph(title, title_style), Spacer(1, 12)]
    for para in content_text.split("\n"):
        clean = para.strip()
        if clean:
            story.append(Paragraph(clean.replace("&", "&amp;"), body_style))
        else:
            story.append(Spacer(1, 8))

    doc.build(story)
    return path


def petition_to_docx(title: str, content_text: str) -> str:
    directory = _ensure_upload_dir()
    filename = f"petition_{uuid.uuid4().hex}.docx"
    path = os.path.join(directory, filename)

    document = Document()
    heading = document.add_heading(title, level=1)
    heading.alignment = 1  # center

    for para in content_text.split("\n"):
        clean = para.strip()
        p = document.add_paragraph(clean)
        p.style.font.size = Pt(11)

    document.save(path)
    return path
